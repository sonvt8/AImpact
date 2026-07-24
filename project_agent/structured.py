import re
import unicodedata
from pathlib import Path
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from rag import NO_EVIDENCE_MESSAGE


DEFAULT_WORKBOOK = Path(__file__).with_name("Phu luc 1.xlsx")
SUMMARY_SHEET = "Tong hop"
DETAIL_SHEET = "3. Ds VHKT"
ROW_ALIASES = {
    "AC": ("ac",),
    "DC": ("dc",),
    "UPS": ("ups",),
    "ĐHCX": ("đhcx", "dhcx", "dieu hoa chinh xac"),
    "UDB/PDU": ("udb pdu", "udb", "pdu"),
}
COLUMN_ALIASES = {
    "UCTT": ("uctt", "ung cuu thong tin"),
    "XLSC": ("xlsc", "xu ly su co"),
    "VHKT": ("vhkt", "van hanh ky thuat"),
}
SITE_ALIASES = {"N4": ("n4",), "N6": ("n6",)}
ERROR_ALIASES = {"Lỗi đơn": ("loi don",), "Lỗi kép": ("loi kep",)}
AGGREGATION_CUES = (
    "bao nhieu",
    "co may",
    "gia tri",
    "so luong",
    "so phuong an",
)


def _normalize(value) -> str:
    value = "" if value is None else str(value).casefold()
    value = "".join(
        character
        for character in unicodedata.normalize("NFD", value)
        if unicodedata.category(character) != "Mn"
    )
    value = value.replace("đ", "d")
    return " ".join(re.findall(r"[0-9a-zđ]+", value))


def _contains(text, phrase) -> bool:
    return f" {phrase} " in f" {text} "


def _has_total(text) -> bool:
    return bool(re.search(r"(?:^| )tong(?= |$)(?! hop(?: |$))", text))


def _has_site(text, site) -> bool:
    return bool(
        re.search(rf"(?<![a-z]){re.escape(_normalize(site))}(?![a-z0-9])", text)
    )


def _claimed_number(text, source, other_source):
    start = text.find(source)
    if start < 0:
        return None
    start += len(source)
    end = text.find(other_source, start)
    match = re.search(
        r"(?<![a-z0-9])\d+(?![a-z0-9])",
        text[start : end if end >= 0 else None],
    )
    return match.group() if match else None


def _claims_summary_detail_conflict(text) -> bool:
    summary = _claimed_number(text, "tong hop", "chi tiet")
    detail = _claimed_number(text, "chi tiet", "tong hop")
    return (
        summary is not None
        and detail is not None
        and summary != detail
        and any(
            _contains(text, cue)
            for cue in ("nhung", "mau thuan", "khong khop", "lech", "so nao")
        )
    )


def _slots(text, aliases) -> list[str]:
    return [
        name
        for name, phrases in aliases.items()
        if any(_contains(text, phrase) for phrase in phrases)
    ]


def _refuse() -> dict:
    return {
        "answer": NO_EVIDENCE_MESSAGE,
        "citations": [],
        "llm_called": False,
    }


def _summary_cell(workbook, row_name, column_name):
    worksheet = workbook[SUMMARY_SHEET]
    header_rows = [
        row
        for row in range(1, min(8, worksheet.max_row) + 1)
        if any(
            _normalize(worksheet.cell(row, column).value) == "su co"
            for column in range(1, worksheet.max_column + 1)
        )
    ]
    if len(header_rows) != 1:
        return None

    header_row = header_rows[0]
    headers = {}
    for cell in worksheet[header_row]:
        headers.setdefault(_normalize(cell.value), []).append(cell.column)
    incident_columns = headers.get("su co", [])
    value_columns = headers.get(_normalize(column_name), [])
    if len(incident_columns) != 1 or len(value_columns) != 1:
        return None

    matching_rows = [
        row
        for row in range(header_row + 1, worksheet.max_row + 1)
        if _normalize(worksheet.cell(row, incident_columns[0]).value)
        == _normalize(row_name)
    ]
    if len(matching_rows) != 1:
        return None
    return worksheet.cell(matching_rows[0], value_columns[0])


def _vhkt_records(workbook, row_name, filename):
    worksheet = workbook[DETAIL_SHEET]
    required_headers = ("stt", "mang", "dau hieu nhan biet", "loai loi")
    header_cells = {
        name: [
            cell
            for row in worksheet.iter_rows(min_row=1, max_row=min(8, worksheet.max_row))
            for cell in row
            if _normalize(cell.value) == name
        ]
        for name in required_headers
    }
    if any(len(cells) != 1 for cells in header_cells.values()):
        return None

    header_row = header_cells["stt"][0].row
    if any(cells[0].row != header_row for cells in header_cells.values()):
        return None
    used_columns = [
        cell.column for cell in worksheet[header_row] if _normalize(cell.value)
    ]
    start_column, end_column = min(used_columns), max(used_columns)
    headers = [
        worksheet.cell(header_row, column).value
        for column in range(start_column, end_column + 1)
    ]

    current_area = None
    records = []
    for row in range(header_row + 1, worksheet.max_row + 1):
        area = worksheet.cell(row, header_cells["mang"][0].column).value
        if area is not None and str(area).strip():
            normalized_area = _normalize(area)
            current_area = next(
                (name for name in ROW_ALIASES if _normalize(name) == normalized_area),
                None,
            )

        stt = worksheet.cell(row, header_cells["stt"][0].column).value
        is_data_row = not isinstance(stt, bool) and bool(
            re.fullmatch(r"\d+(?:\.\d+)?", str(stt or "").strip())
        )
        if not is_data_row:
            if area is None or not str(area).strip():
                current_area = None
            continue
        if current_area != row_name:
            continue

        values = [
            worksheet.cell(row, column).value
            for column in range(start_column, end_column + 1)
        ]
        records.append(
            {
                "site": _normalize(
                    worksheet.cell(
                        row,
                        header_cells["dau hieu nhan biet"][0].column,
                    ).value
                ),
                "error": _normalize(
                    worksheet.cell(row, header_cells["loai loi"][0].column).value
                ),
                "citation": {
                    "filename": filename,
                    "sheet_name": DETAIL_SHEET,
                    "locator": f"{DETAIL_SHEET}!A{row}:L{row}",
                    "content": "\n".join(
                        f"{header}: {'' if value is None else value}"
                        for header, value in zip(headers, values)
                    ),
                },
            }
        )
    return records


def _total_is_consistent(values_workbook, formula_workbook, cell) -> bool:
    values = [
        values_workbook[SUMMARY_SHEET].cell(cell.row, column).value
        for column in range(3, 6)
    ]
    if any(
        value is not None
        and (not isinstance(value, (int, float)) or isinstance(value, bool))
        for value in values
    ):
        return False
    formula = formula_workbook[SUMMARY_SHEET][cell.coordinate].value
    return cell.value == sum(value or 0 for value in values) and (
        isinstance(formula, str)
        and re.sub(r"\s+", "", formula).upper() == f"=SUM(C{cell.row}:E{cell.row})"
    )


def execute(query, workbook_path=DEFAULT_WORKBOOK):
    """Return a deterministic structured answer, refusal, or None for RAG fallback."""
    if not isinstance(query, str) or not query.strip():
        return None

    text = _normalize(query)
    if _claims_summary_detail_conflict(text):
        return _refuse()
    has_total = _has_total(text)
    if not has_total and not any(_contains(text, cue) for cue in AGGREGATION_CUES):
        return None

    rows = _slots(text, ROW_ALIASES)
    columns = _slots(text, COLUMN_ALIASES)
    sites = _slots(text, SITE_ALIASES)
    errors = _slots(text, ERROR_ALIASES)
    if has_total:
        columns.append("TỔNG")
    columns = list(dict.fromkeys(columns))
    if (
        len(rows) != 1
        or len(columns) != 1
        or len(sites) > 1
        or len(errors) > 1
        or (_contains(text, "loi") and not errors)
        or ((sites or errors) and columns != ["VHKT"])
    ):
        return _refuse()

    workbook = formula_workbook = None
    try:
        workbook = openpyxl.load_workbook(
            workbook_path,
            data_only=True,
            read_only=True,
        )
        if columns[0] == "TỔNG":
            formula_workbook = openpyxl.load_workbook(
                workbook_path,
                data_only=False,
                read_only=True,
            )
    except (OSError, BadZipFile, InvalidFileException):
        if workbook is not None:
            workbook.close()
        return _refuse()

    try:
        cell = _summary_cell(workbook, rows[0], columns[0])
        if (
            cell is None
            or not isinstance(cell.value, (int, float))
            or isinstance(cell.value, bool)
        ):
            return _refuse()
        if columns[0] == "TỔNG" and not _total_is_consistent(
            workbook,
            formula_workbook,
            cell,
        ):
            return _refuse()
        if columns[0] == "VHKT":
            records = _vhkt_records(workbook, rows[0], Path(workbook_path).name)
            if records is None or len(records) != cell.value:
                return _refuse()
            if sites or errors:
                records = [
                    record
                    for record in records
                    if (not sites or _has_site(record["site"], sites[0]))
                    and (not errors or record["error"] == _normalize(errors[0]))
                ]
                if not records:
                    return _refuse()
                return {
                    "answer": len(records),
                    "citations": [record["citation"] for record in records],
                    "llm_called": False,
                }
        return {
            "answer": cell.value,
            "citations": [
                {
                    "filename": Path(workbook_path).name,
                    "sheet_name": SUMMARY_SHEET,
                    "locator": f"{SUMMARY_SHEET}!{cell.coordinate}",
                    "content": str(cell.value),
                }
            ],
            "llm_called": False,
        }
    except KeyError:
        return _refuse()
    finally:
        workbook.close()
        if formula_workbook is not None:
            formula_workbook.close()
