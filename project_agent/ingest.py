import csv
import re
from dataclasses import dataclass

import openpyxl
from openpyxl.utils import get_column_letter


@dataclass
class Record:
    source_type: str
    locator: str
    text_verbatim: str
    sheet_name: str
    row_index: int | None
    headers: list[str]
    fields: dict[str, object]
    section_path: list[str]
    stt: object | None
    is_section: bool


def parse_file(path) -> list[Record]:
    extension = str(path).lower().rsplit(".", 1)[-1]
    if extension == "pdf":
        return parse_pdf(path)
    if extension == "docx":
        return parse_docx(path)
    if extension == "xlsx":
        return parse_xlsx(path)
    if extension == "txt":
        return parse_txt(path)
    if extension == "csv":
        return parse_csv(path)
    raise ValueError(f"Unsupported file type: {path}")


def parse_pdf(path) -> list[Record]:
    import pdfplumber

    records = []
    with pdfplumber.open(path) as document:
        for page_number, page in enumerate(document.pages, start=1):
            table_records = []
            try:
                tables = page.extract_tables() or []
            except Exception:
                tables = []
            for table_index, table in enumerate(tables, start=1):
                if not table:
                    continue
                width = max((len(row or []) for row in table), default=0)
                if not width:
                    continue
                first_row = next(
                    (row for row in table if row and any(value not in (None, "") for value in row)),
                    [],
                )
                headers = _unique_headers(list(first_row) + [None] * (width - len(first_row)), 1)
                for row_index, row in enumerate(table, start=1):
                    values = list(row or [])[:width] + [None] * max(0, width - len(row or []))
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    normalized = ["" if value is None else value for value in values]
                    table_records.append(
                        Record(
                            source_type="pdf",
                            locator=f"p.{page_number}#t{table_index}-r{row_index}",
                            text_verbatim="\t".join(str(value) for value in normalized),
                            sheet_name="",
                            row_index=row_index,
                            headers=list(headers),
                            fields=dict(zip(headers, normalized)),
                            section_path=[],
                            stt=normalized[0] if normalized else None,
                            is_section=False,
                        )
                    )
            if table_records:
                records.extend(table_records)
                try:
                    table_bboxes = [table.bbox for table in page.find_tables()]
                    outside_tables = page.filter(
                        lambda obj: obj.get("object_type") != "char"
                        or not any(
                            x0 <= (obj["x0"] + obj["x1"]) / 2 <= x1
                            and top <= (obj["top"] + obj["bottom"]) / 2 <= bottom
                            for x0, top, x1, bottom in table_bboxes
                        )
                    )
                    text = outside_tables.extract_text() or ""
                except Exception:
                    text = page.extract_text() or ""
            else:
                text = page.extract_text() or ""
            if text.strip():
                records.append(
                    Record(
                        source_type="pdf",
                        locator=f"p.{page_number}#text" if table_records else f"p.{page_number}",
                        text_verbatim=text,
                        sheet_name="",
                        row_index=page_number,
                        headers=[],
                        fields={},
                        section_path=[],
                        stt=None,
                        is_section=False,
                    )
                )
    return records


def parse_docx(path) -> list[Record]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(path)
    records = []
    section_path = []
    paragraph_index = 0
    table_index = 0

    for child in document.element.body.iterchildren():
        kind = child.tag.rsplit("}", 1)[-1]
        if kind == "p":
            paragraph_index += 1
            paragraph = Paragraph(child, document)
            text = paragraph.text
            if not text.strip():
                continue
            style = paragraph.style.style_id if paragraph.style else ""
            is_heading = style.lower().startswith("heading")
            if is_heading:
                match = re.search(r"(\d+)$", style)
                level = int(match.group(1)) if match else 1
                section_path = section_path[: level - 1] + [text]
            records.append(
                Record(
                    source_type="docx",
                    locator=f"para:{paragraph_index}",
                    text_verbatim=text,
                    sheet_name="",
                    row_index=paragraph_index,
                    headers=[],
                    fields={},
                    section_path=list(section_path),
                    stt=None,
                    is_section=False,
                )
            )
        elif kind == "tbl":
            table_index += 1
            table = Table(child, document)
            for row_index, row in enumerate(table.rows, start=1):
                for column_index, cell in enumerate(row.cells, start=1):
                    if not cell.text.strip():
                        continue
                    records.append(
                        Record(
                            source_type="docx",
                            locator=f"table:{table_index},r{row_index},c{column_index}",
                            text_verbatim=cell.text,
                            sheet_name="",
                            row_index=row_index,
                            headers=[],
                            fields={},
                            section_path=list(section_path),
                            stt=None,
                            is_section=False,
                        )
                    )
    return records


def parse_xlsx(path) -> list[Record]:
    workbook = openpyxl.load_workbook(path, data_only=False)
    records = []
    try:
        for worksheet in workbook.worksheets:
            records.extend(_parse_worksheet(worksheet))
    finally:
        workbook.close()
    return records


def parse_txt(path) -> list[Record]:
    with open(path, "r", encoding="utf-8", newline="") as source:
        text = source.read()
    return [
        Record(
            source_type="txt",
            locator=f"chars:0-{len(text)}",
            text_verbatim=text,
            sheet_name="",
            row_index=None,
            headers=[],
            fields={},
            section_path=[],
            stt=None,
            is_section=False,
        )
    ]


def parse_csv(path) -> list[Record]:
    with open(path, "r", encoding="utf-8-sig", newline="") as source:
        rows = list(csv.reader(source))
        if not rows:
            return []
        header_row = rows[0]
        width = max(len(row) for row in rows)
        header_row += [None] * (width - len(header_row))
        headers = _unique_headers(header_row, 1)
        records = []
        for row_index, row in enumerate(rows[1:], start=2):
            values = row[: len(headers)] + [""] * max(0, len(headers) - len(row))
            fields = dict(zip(headers, values))
            records.append(
                Record(
                    source_type="csv",
                    locator=f"row:{row_index}",
                    text_verbatim=_record_text(headers, values, []),
                    sheet_name="",
                    row_index=row_index,
                    headers=list(headers),
                    fields=fields,
                    section_path=[],
                    stt=values[0] if values else None,
                    is_section=False,
                )
            )
    return records


def _parse_worksheet(worksheet) -> list[Record]:
    merged_values = _merged_values(worksheet)
    blocks = _row_blocks(worksheet, merged_values)
    if not blocks:
        return []

    detected = [
        (start_row, end_row, _find_header_row(worksheet, start_row, end_row))
        for start_row, end_row in blocks
    ]
    if not any(header_row is not None for _, _, header_row in detected):
        return [_free_text_record(worksheet, blocks[0][0], blocks[-1][1])]

    records = []
    for start_row, end_row, header_row in detected:
        if header_row is None:
            records.append(_free_text_record(worksheet, start_row, end_row))
        else:
            records.extend(_parse_table_block(worksheet, header_row, end_row, merged_values))
    return records


def _parse_table_block(worksheet, header_row, end_row, merged_values) -> list[Record]:
    used_columns = [
        column
        for column in range(1, worksheet.max_column + 1)
        if any(
            value is not None and str(value).strip()
            for value in (
                _cell_value(worksheet, row, column, merged_values)
                for row in range(header_row, end_row + 1)
            )
        )
    ]
    if not used_columns:
        return []

    start_column, end_column = min(used_columns), max(used_columns)
    header_values = [
        _cell_value(worksheet, header_row, column, merged_values)
        for column in range(start_column, end_column + 1)
    ]
    headers = _unique_headers(header_values, start_column)
    section_path = []
    records = []

    for row_index in range(header_row + 1, end_row + 1):
        values = [
            _cell_value(worksheet, row_index, column, merged_values)
            for column in range(start_column, end_column + 1)
        ]
        if not any(value is not None and str(value).strip() for value in values):
            continue

        nonnull = [(index, value) for index, value in enumerate(values) if value is not None]
        distinct_nonfirst = {
            str(value).strip() for index, value in nonnull if index != 0
        }
        is_section = len(values) > 2 and len(distinct_nonfirst) <= 1
        fields = {
            header: "" if value is None else value
            for header, value in zip(headers, values)
        }
        locator = (
            f"{worksheet.title}!{get_column_letter(start_column)}{row_index}:"
            f"{get_column_letter(end_column)}{row_index}"
        )

        if is_section:
            label = _section_label(values)
            level = _section_level(values[0])
            section_path = section_path[:level] + [label]
            text_verbatim = label
        else:
            text_verbatim = _record_text(headers, values, section_path)

        records.append(
            Record(
                source_type="xlsx",
                locator=locator,
                text_verbatim=text_verbatim,
                sheet_name=worksheet.title,
                row_index=row_index,
                headers=list(headers),
                fields=fields,
                section_path=list(section_path),
                stt=values[0] if values else None,
                is_section=is_section,
            )
        )
    return records


def _row_blocks(worksheet, merged_values) -> list[tuple[int, int]]:
    blocks = []
    start_row = None
    for row in range(1, worksheet.max_row + 1):
        has_content = any(
            value is not None and str(value).strip()
            for value in (
                _cell_value(worksheet, row, column, merged_values)
                for column in range(1, worksheet.max_column + 1)
            )
        )
        if has_content and start_row is None:
            start_row = row
        elif not has_content and start_row is not None:
            blocks.append((start_row, row - 1))
            start_row = None
    if start_row is not None:
        blocks.append((start_row, worksheet.max_row))
    return blocks


def _find_header_row(worksheet, start_row=1, end_row=None) -> int | None:
    end_row = worksheet.max_row if end_row is None else end_row
    first_row_values = [
        worksheet.cell(start_row, column).value
        for column in range(1, worksheet.max_column + 1)
    ]
    first_row_nonempty = [
        value for value in first_row_values if value is not None and str(value).strip()
    ]
    block_nonempty = [
        worksheet.cell(row, column).value
        for row in range(start_row, end_row + 1)
        for column in range(1, worksheet.max_column + 1)
        if worksheet.cell(row, column).value is not None
        and str(worksheet.cell(row, column).value).strip()
    ]
    if (
        start_row < end_row
        and len(first_row_nonempty) >= 2
        and all(isinstance(value, str) for value in block_nonempty)
    ):
        return start_row

    candidates = []
    for row in range(start_row, end_row):
        values = [worksheet.cell(row, column).value for column in range(1, worksheet.max_column + 1)]
        occupied = [index for index, value in enumerate(values) if value is not None and str(value).strip()]
        strings = sum(
            isinstance(values[index], str) and bool(values[index].strip())
            for index in occupied
        )
        if strings < 2:
            continue

        start_column, end_column = occupied[0], occupied[-1]
        width = end_column - start_column + 1
        string_density = strings / width
        if string_density < 0.5:
            continue

        data_rows = []
        for data_row in range(row + 1, min(end_row, row + 12) + 1):
            data_values = [
                worksheet.cell(data_row, column + 1).value
                for column in range(start_column, end_column + 1)
            ]
            if any(value is not None and str(value).strip() for value in data_values):
                data_rows.append(data_values)
        if not data_rows:
            continue

        data_density = sum(
            sum(value is not None and bool(str(value).strip()) for value in data_values) / width
            for data_values in data_rows
        ) / len(data_rows)
        if data_density < 0.5:
            continue

        consistencies = []
        for column in range(width):
            kinds = [_value_kind(data_values[column]) for data_values in data_rows if data_values[column] not in (None, "")]
            if kinds:
                consistencies.append(max(kinds.count(kind) for kind in set(kinds)) / len(kinds))
        consistency = sum(consistencies) / len(consistencies) if consistencies else 0
        candidates.append((strings + string_density + data_density + consistency, row))

    return max(candidates, default=(0, None))[1]


def _value_kind(value):
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, (int, float)):
        return "number"
    return type(value)


def _free_text_record(worksheet, start_row, end_row) -> Record:
    used_columns = [
        column
        for column in range(1, worksheet.max_column + 1)
        if any(
            worksheet.cell(row, column).value is not None
            and str(worksheet.cell(row, column).value).strip()
            for row in range(start_row, end_row + 1)
        )
    ]
    start_column, end_column = min(used_columns), max(used_columns)
    lines = []
    for row in range(start_row, end_row + 1):
        values = [worksheet.cell(row, column).value for column in range(start_column, end_column + 1)]
        lines.append("\t".join("" if value is None else str(value) for value in values).rstrip("\t"))
    return Record(
        source_type="xlsx",
        locator=(
            f"{worksheet.title}!{get_column_letter(start_column)}{start_row}:"
            f"{get_column_letter(end_column)}{end_row}"
        ),
        text_verbatim="\n".join(lines).strip("\n"),
        sheet_name=worksheet.title,
        row_index=start_row,
        headers=[],
        fields={},
        section_path=[],
        stt=None,
        is_section=False,
    )


def _merged_values(worksheet) -> dict[tuple[int, int], object]:
    values = {}
    for merged_range in worksheet.merged_cells.ranges:
        anchor = worksheet.cell(merged_range.min_row, merged_range.min_col).value
        for row in range(merged_range.min_row, merged_range.max_row + 1):
            for column in range(merged_range.min_col, merged_range.max_col + 1):
                values[(row, column)] = anchor
    return values


def _cell_value(worksheet, row: int, column: int, merged_values):
    cell = worksheet.cell(row, column)
    if cell.value is None and (row, column) in merged_values:
        return merged_values[(row, column)]
    if cell.value is None and cell.data_type == "inlineStr":
        return ""
    return cell.value


def _unique_headers(values, start_column: int) -> list[str]:
    counts = {}
    headers = []
    for offset, value in enumerate(values):
        base = str(value) if value is not None and str(value).strip() else get_column_letter(start_column + offset)
        counts[base] = counts.get(base, 0) + 1
        headers.append(base if counts[base] == 1 else f"{base}_{counts[base]}")
    return headers


def _section_label(values) -> str:
    parts = []
    seen = set()
    for value in values:
        if value is None or not str(value).strip():
            continue
        key = str(value).strip()
        if key not in seen:
            seen.add(key)
            parts.append(str(value))
    return " | ".join(parts)


def _section_level(value) -> int:
    outline = "" if value is None else str(value).strip()
    if re.fullmatch(r"[IVXLC]+", outline, re.IGNORECASE):
        return 0
    if re.fullmatch(r"[IVXLC0-9]+(?:\.\d+)+", outline, re.IGNORECASE):
        return outline.count(".")
    return 0


def _record_text(headers, values, section_path) -> str:
    body = "\n".join(
        f"{header}: {'' if value is None else value}"
        for header, value in zip(headers, values)
    )
    return "\n".join([*section_path, body]) if section_path else body
