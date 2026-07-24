import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from scripts import eval_retrieval


def test_metrics_use_only_gated_citations(monkeypatch):
    hits = [
        {
            "text_verbatim": "first",
            "metadata": {"locator": "Sheet!A1:C1"},
            "similarity": 0.9,
            "lexical_gate": True,
        },
        {
            "text_verbatim": "second",
            "metadata": {"locator": "Sheet!A2:C2"},
            "similarity": 0.8,
            "lexical_gate": True,
        },
    ]
    monkeypatch.setattr(
        eval_retrieval,
        "_query_timed",
        lambda *args: (
            hits,
            hits[:1],
            {"embed_ms": 1.0, "chroma_ms": 2.0, "gate_ms": 3.0},
        ),
    )

    row = eval_retrieval.evaluate(
        object(),
        object(),
        [
            {
                "id": "q1",
                "route": "vector",
                "query": "query",
                "acceptable_locators": ["Sheet!A2:C2"],
            }
        ],
        top_k=5,
        threshold=0.5,
        doc_path=Path("unused.xlsx"),
    )[0]

    assert row["actual_route"] == "vector"
    assert row["route_correct"] is True
    assert row["locator_recall_5"] is True
    assert row["citation_precision"] == 0.0
    assert row["citation_recall"] == 0.0
    assert row["relevant_survives_gate"] is False


def test_auto_selects_fasttext_file_and_onnx_directory(monkeypatch, tmp_path):
    class FakeFastText:
        def __init__(self, **kwargs):
            self.kwargs = kwargs

    class FakeOnnx:
        def __init__(self, path, **kwargs):
            self.path = path
            self.kwargs = kwargs

    monkeypatch.setattr(eval_retrieval.embedding, "FastTextEmbedder", FakeFastText)
    monkeypatch.setattr(eval_retrieval.embedding, "LocalOnnxEmbedder", FakeOnnx)
    fasttext_path = tmp_path / "model.bin"
    fasttext_path.touch()
    onnx_path = tmp_path / "bundle"
    onnx_path.mkdir()

    fasttext, fasttext_kind = eval_retrieval.select_embedder("auto", fasttext_path)
    onnx, onnx_kind = eval_retrieval.select_embedder("auto", onnx_path)

    assert fasttext_kind == "fasttext"
    assert fasttext.kwargs["model_path"] == str(fasttext_path)
    assert onnx_kind == "onnx-e5"
    assert onnx.path == onnx_path
    assert onnx.kwargs == {"max_length": 96, "batch_size": 16}


def test_structured_and_refusal_routes_record_real_metrics(monkeypatch):
    results = {
        "structured": {
            "answer": 6,
            "citations": [{"locator": "Tong hop!E3"}],
            "llm_called": False,
        },
        "refuse": {
            "answer": eval_retrieval.rag.NO_EVIDENCE_MESSAGE,
            "citations": [],
            "llm_called": False,
        },
    }
    monkeypatch.setattr(
        eval_retrieval.structured,
        "execute",
        lambda query, path: results[query],
    )
    monkeypatch.setattr(
        eval_retrieval,
        "_query_timed",
        lambda *args: pytest.fail("structured refusal must not query vector index"),
    )

    rows = eval_retrieval.evaluate(
        object(),
        object(),
        [
            {
                "id": "s1",
                "route": "structured",
                "query": "structured",
                "acceptable_locators": ["Tong hop!E3"],
                "expected_answer": 6,
            },
            {
                "id": "r1",
                "route": "refuse",
                "query": "refuse",
                "acceptable_locators": [],
            },
        ],
        top_k=5,
        threshold=0.5,
        doc_path="workbook.xlsx",
    )

    assert rows[0]["actual_route"] == "structured"
    assert rows[0]["route_correct"] is True
    assert rows[0]["answer_correct"] is True
    assert rows[0]["citation_precision"] == 1.0
    assert rows[0]["citation_recall"] == 1.0
    assert rows[1]["actual_route"] == "refuse"
    assert rows[1]["route_correct"] is True
    assert eval_retrieval.refusal_metrics(rows)["recall"] == 1.0


def test_golden_expected_values_match_workbook():
    root = Path(__file__).parents[1]
    golden = json.loads(
        (root / "scripts" / "golden_queries.json").read_text(encoding="utf-8")
    )
    workbook = load_workbook(root / golden["document"], data_only=True, read_only=True)
    try:
        for item in golden["queries"]:
            for locator in item["acceptable_locators"]:
                sheet, cell_range = locator.rsplit("!", 1)
                min_col, min_row, max_col, max_row = range_boundaries(cell_range)
                values = [
                    workbook[sheet].cell(row, column).value
                    for row in range(min_row, max_row + 1)
                    for column in range(min_col, max_col + 1)
                ]
                assert any(value is not None for value in values), locator

            if item["route"] != "structured":
                continue
            if item["expected_operation"] == "count":
                assert item["expected_answer"] == len(item["acceptable_locators"])
                continue
            sheet, cell = item["acceptable_locators"][0].rsplit("!", 1)
            assert workbook[sheet][cell].value == item["expected_answer"]
    finally:
        workbook.close()
