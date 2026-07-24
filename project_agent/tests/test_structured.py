import re
import shutil
from pathlib import Path
from zipfile import ZipFile

import openpyxl
import pytest

from rag import NO_EVIDENCE_MESSAGE
from structured import execute


WORKBOOK = Path(__file__).resolve().parents[1] / "Phu luc 1.xlsx"
REFUSAL = {
    "answer": NO_EVIDENCE_MESSAGE,
    "citations": [],
    "llm_called": False,
}


def _replace_cached_value(path, coordinate, value):
    worksheet_path = "xl/worksheets/sheet1.xml"
    with ZipFile(path) as source:
        entries = [(info, source.read(info.filename)) for info in source.infolist()]
    for index, (info, content) in enumerate(entries):
        if info.filename != worksheet_path:
            continue
        content, replacements = re.subn(
            rb'(<c r="' + coordinate.encode() + rb'"[^>]*>.*?<v>)[^<]*(</v>)',
            rb"\g<1>" + str(value).encode() + rb"\2",
            content,
            count=1,
        )
        assert replacements == 1
        entries[index] = (info, content)
        break
    with ZipFile(path, "w") as target:
        for info, content in entries:
            target.writestr(info, content)


def test_ac_vhkt_is_six_with_verbatim_cell_citation():
    result = execute("Có bao nhiêu sự cố AC trong VHKT?", WORKBOOK)

    assert result == {
        "answer": 6,
        "citations": [
            {
                "filename": "Phu luc 1.xlsx",
                "sheet_name": "Tong hop",
                "locator": "Tong hop!E3",
                "content": "6",
            }
        ],
        "llm_called": False,
    }


def test_summary_value_lookup():
    result = execute("Số phương án UCTT của sự cố AC là bao nhiêu?", WORKBOOK)

    assert result["answer"] == 28
    assert result["citations"][0]["locator"] == "Tong hop!C3"
    assert result["citations"][0]["content"] == "28"


def test_vhkt_multi_filter_with_forward_filled_area(tmp_path):
    path = tmp_path / WORKBOOK.name
    shutil.copy2(WORKBOOK, path)
    workbook = openpyxl.load_workbook(path)
    for row in range(6, 11):
        workbook["3. Ds VHKT"].cell(row, 2).value = None
    workbook.save(path)
    workbook.close()

    result = execute(
        "Trong danh sách VHKT, có bao nhiêu tình huống AC cho N4 và là lỗi đơn?",
        path,
    )

    assert result["answer"] == 2
    assert [citation["locator"] for citation in result["citations"]] == [
        "3. Ds VHKT!A8:L8",
        "3. Ds VHKT!A9:L9",
    ]
    assert all(
        "N4" in citation["content"] and "Lỗi đơn" in citation["content"]
        for citation in result["citations"]
    )


def test_refuses_stale_total_formula_cache(tmp_path):
    query = "Tổng cộng các phương án AC là bao nhiêu?"
    assert execute(query, WORKBOOK)["answer"] == 61

    path = tmp_path / WORKBOOK.name
    shutil.copy2(WORKBOOK, path)
    _replace_cached_value(path, "F3", 60)

    assert execute(query, path) == REFUSAL


@pytest.mark.parametrize(
    "query",
    [
        "Có bao nhiêu sự cố AC?",
        "Có bao nhiêu sự cố AC và DC trong VHKT?",
        "AC trong UCTT và XLSC có bao nhiêu phương án?",
        "Có bao nhiêu sự cố XYZ trong VHKT?",
    ],
)
def test_refuses_missing_or_ambiguous_slots(query):
    assert execute(query, WORKBOOK) == REFUSAL


def test_refuses_summary_detail_conflict(tmp_path):
    conflict_path = tmp_path / WORKBOOK.name
    shutil.copy2(WORKBOOK, conflict_path)
    workbook = openpyxl.load_workbook(conflict_path)
    workbook["3. Ds VHKT"]["B10"] = "DC"
    workbook.save(conflict_path)
    workbook.close()

    assert execute("Có bao nhiêu sự cố AC trong VHKT?", conflict_path) == REFUSAL


def test_non_structured_question_falls_back():
    assert execute("N6 mất lộ điện nổi thì xử lý thế nào?", WORKBOOK) is None


def test_refuses_claimed_summary_detail_conflict_before_rag_fallback():
    query = (
        "Nếu sheet Tổng hợp ghi AC trong VHKT là 6 nhưng danh sách chi tiết "
        "chỉ có 5 thì phải dùng số nào?"
    )

    assert execute(query, WORKBOOK) == REFUSAL
    assert execute(
        "Sheet Tổng hợp và danh sách chi tiết mô tả N4 khác nhau thì xử lý thế nào?",
        WORKBOOK,
    ) is None
