"""C-GEN-1 acceptance harness — generic parser robustness.

Test-driven contract for hardening ingest.parse_file to handle DIVERSE,
user-supplied Excel / PDF / Docx structures (tables + quick-reference layouts).
These tests are written against BEHAVIOUR, not implementation, so Codex has
freedom in HOW to satisfy them. Generic-first: NO fixture uses the example
workbook's sheet names or terminology.

Run:  pytest tests/test_ingest_robustness.py -v
Contract for Codex:
  * Make every test here pass.
  * Do NOT break the existing 46 tests (run full suite).
  * Do NOT change gate/citation/verbatim semantics elsewhere.
  * Keep parsing generic — no hard-coded headers/sheet names/terms.

Fixtures are generated at runtime (openpyxl / python-docx / reportlab) so the
harness is self-contained and portable.
"""

import os
import sys

import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import ingest  # noqa: E402


def _data_records(path):
    recs = ingest.parse_file(str(path))
    return [r for r in recs if not r.is_section]


# ----------------------------------------------------------------------------
# EXCEL
# ----------------------------------------------------------------------------
def _xlsx(path, sheets):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets:
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    wb.save(str(path))


def test_regression_offset_header_within_window(tmp_path):
    """Title/metadata rows above the header (still within scan window) must not
    break header detection. This already works; guard against regression."""
    p = tmp_path / "offset.xlsx"
    _xlsx(p, [("Q3", [
        ["QUARTERLY REPORT"], ["Confidential"], [],
        ["Region", "Product", "Units", "Revenue"],
        ["North", "Widget", 100, 5000],
        ["South", "Gadget", 80, 4800],
    ])])
    data = _data_records(p)
    joined = "\n".join(r.text_verbatim for r in data)
    assert "Region: North" in joined
    assert "Product: Widget" in joined
    assert len([r for r in data if "Region:" in r.text_verbatim]) == 2


def test_header_beyond_scan_window(tmp_path):
    """TARGET: header pushed past the first rows (long preamble) must still be
    detected, not mistaken for data. Current parser fails this."""
    p = tmp_path / "late.xlsx"
    rows = [[f"preamble line {i + 1}"] for i in range(9)]
    rows += [["Code", "Description", "Qty"], ["A1", "Bolt", 10], ["A2", "Nut", 20]]
    _xlsx(p, [("Late", rows)])
    data = _data_records(p)
    joined = "\n".join(r.text_verbatim for r in data)
    assert "Description: Bolt" in joined, "header row not detected past window"
    assert "Code: A1" in joined
    assert "Qty: 10" in joined
    # preamble lines must not be misread as field headers
    assert "preamble line 1:" not in joined


def test_multiple_tables_one_sheet(tmp_path):
    """TARGET: two independent tables in one sheet (separated by a blank row,
    different headers) — rows from BOTH must be captured with their OWN headers.
    Current single-header-per-sheet assumption fails this."""
    p = tmp_path / "multi.xlsx"
    rows = [
        ["Region", "Product", "Units"],
        ["North", "Widget", 100],
        ["South", "Gadget", 80],
        [],
        [],
        ["SKU", "Item", "Stock"],
        ["SKU-01", "Cable", 340],
        ["SKU-02", "Router", 12],
    ]
    _xlsx(p, [("Mixed", rows)])
    data = _data_records(p)
    joined = "\n".join(r.text_verbatim for r in data)
    assert "Region: North" in joined and "Product: Widget" in joined
    assert "SKU: SKU-01" in joined and "Item: Cable" in joined
    assert "Stock: 340" in joined


# ----------------------------------------------------------------------------
# DOCX
# ----------------------------------------------------------------------------
def test_docx_granularity_and_table(tmp_path):
    """TARGET: a docx with multiple headings + a table must yield MORE THAN ONE
    record with distinct locators (not one whole-file blob), and table cell
    values must be present. Current parser returns a single record."""
    import docx

    p = tmp_path / "guide.docx"
    doc = docx.Document()
    doc.add_heading("Section A", level=1)
    doc.add_paragraph("Steps for section A.")
    doc.add_heading("Section B", level=1)
    doc.add_paragraph("Steps for section B.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Key"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Voltage"
    table.cell(1, 1).text = "220V"
    doc.save(str(p))
    data = _data_records(p)
    assert len(data) >= 2, "docx not segmented into multiple records"
    locators = {r.locator for r in data}
    assert len(locators) >= 2, "docx records must have distinct locators"
    joined = "\n".join(r.text_verbatim for r in data)
    assert "Section A" in joined and "Section B" in joined
    assert "Voltage" in joined and "220V" in joined


# ----------------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------------
def test_pdf_table_structure(tmp_path):
    """TARGET: a PDF containing a bordered table must preserve the ROW grouping
    (each data row's cells kept together), not flatten the whole page into one
    blob. ASCII content avoids font-diacritic confounds in the fixture."""
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    p = tmp_path / "table.pdf"
    rows = [
        ["Code", "Problem", "Action"],
        ["INC-01", "Power loss", "Switch feed"],
        ["INC-02", "Battery drain", "Start generator"],
    ]
    t = Table(rows)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
    SimpleDocTemplate(str(p), pagesize=A4).build(
        [Paragraph("Quick reference", getSampleStyleSheet()["Title"]), t]
    )
    data = _data_records(p)
    # A row's cells must co-occur in a record that does NOT also contain the
    # OTHER row's unique marker — i.e. rows are segmented, not flattened into
    # one page blob. A single whole-page record fails this (both markers present).
    def row_isolated(cells, other_marker):
        return any(
            all(c in r.text_verbatim for c in cells)
            and other_marker not in r.text_verbatim
            for r in data
        )

    assert row_isolated(["INC-01", "Power loss", "Switch feed"], "INC-02"), \
        "PDF table rows not segmented (page flattened into one blob)"
    assert row_isolated(["INC-02", "Battery drain", "Start generator"], "INC-01")


# ----------------------------------------------------------------------------
# REGRESSION PIN — example workbook still parses to the known-good count.
# Skips automatically if the file is not alongside the tests.
# ----------------------------------------------------------------------------
def test_regression_example_workbook_count():
    candidates = [
        os.path.join(os.path.dirname(__file__), "..", "Phu luc 1.xlsx"),
        "Phu luc 1.xlsx",
    ]
    path = next((c for c in candidates if os.path.exists(c)), None)
    if path is None:
        pytest.skip("example workbook not present")
    data = _data_records(path)
    assert len(data) == 357, f"example workbook regressed: {len(data)} != 357"


# ----------------------------------------------------------------------------
# C-GEN-1-FIX additions — defects found by adversarial probing of the first pass.
# ----------------------------------------------------------------------------
def test_docx_heading_is_indexable(tmp_path):
    """DEFECT #1: docx headings must remain INDEXABLE. index.py drops is_section
    records before embedding, so heading text marked is_section would become
    unsearchable. Heading content must survive the non-section filter."""
    import docx

    p = tmp_path / "headings.docx"
    d = docx.Document()
    d.add_heading("Bảng tra cứu nhanh sự cố điện", level=1)
    d.add_paragraph("Nội dung mô tả.")
    d.save(str(p))
    recs = ingest.parse_file(str(p))
    indexable = [r for r in recs if not r.is_section]  # what index.py keeps
    joined = "\n".join(r.text_verbatim for r in indexable)
    assert "Bảng tra cứu nhanh sự cố điện" in joined, \
        "heading dropped from indexable records (is_section) -> unsearchable"


def test_pdf_table_and_surrounding_text(tmp_path):
    """DEFECT #2: a PDF page with BOTH a paragraph and a table must keep the
    non-table text, not only the table rows."""
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib import colors

    p = tmp_path / "mixed.pdf"
    t = Table([["Code", "Action"], ["X1", "Do thing"]])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.black)]))
    SimpleDocTemplate(str(p), pagesize=A4).build(
        [Paragraph("Introductory guidance text on this page", getSampleStyleSheet()["Normal"]), t]
    )
    data = _data_records(p)
    joined = "\n".join(r.text_verbatim for r in data)
    assert "Introductory guidance text on this page" in joined, \
        "non-table text on a table page was dropped"
    # table rows still segmented
    assert any("X1" in r.text_verbatim and "Do thing" in r.text_verbatim for r in data)


def test_all_string_table_first_row_is_header(tmp_path):
    """DEFECT #3: when header and data are indistinguishable by type (all-string,
    e.g. a quick-reference table), the FIRST row of a block must be treated as the
    header rather than the scorer picking an arbitrary interior row."""
    p = tmp_path / "allstring.xlsx"
    _xlsx(p, [("Ref", [
        ["Triệu chứng", "Nguyên nhân", "Xử lý"],
        ["Mất điện", "Trip MCCB", "Đóng lại MCCB"],
        ["UPS kêu", "Ắc quy yếu", "Thay ắc quy"],
    ])])
    data = _data_records(p)
    joined = "\n".join(r.text_verbatim for r in data)
    assert "Triệu chứng: Mất điện" in joined, "first row not used as header for all-string table"
    assert "Xử lý: Đóng lại MCCB" in joined
    assert len([r for r in data if "Triệu chứng:" in r.text_verbatim]) == 2


def test_three_tables_all_string(tmp_path):
    """DEFECT #3 (extended): 3+ tables of all-string cells in one sheet — every
    block's first row is its header; all data rows captured."""
    p = tmp_path / "three.xlsx"
    _xlsx(p, [("S", [
        ["A", "B"], ["a1", "b1"], [], [],
        ["C", "D"], ["c1", "d1"], [], [],
        ["E", "F"], ["e1", "f1"], ["e2", "f2"],
    ])])
    data = _data_records(p)
    joined = "\n".join(r.text_verbatim for r in data)
    for expect in ["A: a1", "C: c1", "E: e1", "E: e2"]:
        assert expect in joined, f"missing {expect!r} (header mis-detected in a later block)"