import openpyxl


def read_table(path, sheet_name) -> list[dict]:
    workbook = openpyxl.load_workbook(path, data_only=True)
    try:
        worksheet = workbook[sheet_name]
        header_row = max(
            range(1, min(8, worksheet.max_row) + 1),
            key=lambda row: sum(
                isinstance(worksheet.cell(row, column).value, str)
                and bool(worksheet.cell(row, column).value.strip())
                for column in range(1, worksheet.max_column + 1)
            ),
        )
        used_columns = [
            column
            for column in range(1, worksheet.max_column + 1)
            if any(
                value is not None and str(value).strip()
                for value in (
                    worksheet.cell(row, column).value
                    for row in range(header_row, worksheet.max_row + 1)
                )
            )
        ]
        if not used_columns:
            return []

        start_column, end_column = min(used_columns), max(used_columns)
        headers = [
            worksheet.cell(header_row, column).value
            for column in range(start_column, end_column + 1)
        ]
        rows = []
        for values in worksheet.iter_rows(
            min_row=header_row + 1,
            min_col=start_column,
            max_col=end_column,
            values_only=True,
        ):
            if any(value is not None and str(value).strip() for value in values):
                rows.append(dict(zip(headers, values)))
        return rows
    finally:
        workbook.close()


def column_totals(rows, columns=None) -> dict:
    if columns is None:
        columns = dict.fromkeys(column for row in rows for column in row)

    totals = {}
    for column in columns:
        values = [
            row.get(column)
            for row in rows
            if isinstance(row.get(column), (int, float))
            and not isinstance(row.get(column), bool)
        ]
        if values:
            totals[column] = sum(values)
    return totals
